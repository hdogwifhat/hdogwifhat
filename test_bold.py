import docx

doc = docx.Document()
p = doc.add_paragraph("כמות האופציות שאנו מחזיקים עבורך בנאמנות עומדת על סך של [סך הכול אופציות]  מתוכן סך של [אופציות מובשלות]  אופציות בשלות וסך של [אופציות לא מובשלות]  אופציות שטרם הגיע מועד הבשלתן,")

replacements = {
    "[סך הכול אופציות] אופציות": "50 אופציות",
    "[סך הכול אופציות]": "50",
    "[אופציות מובשלות]  אופציות בשלות": "25 אופציות בשלות",
    "[אופציות לא מובשלות]  אופציות שטרם הגיע מועד הבשלתן": "25 אופציות שטרם הגיע מועד הבשלתן",
}

bold_keys = [
    "[סך הכול אופציות] אופציות",
    "[אופציות מובשלות]  אופציות בשלות",
    "[אופציות לא מובשלות]  אופציות שטרם הגיע מועד הבשלתן"
]

def advanced_replace(paragraphs, replacements, bold_keys):
    for p in paragraphs:
        text = p.text
        # Sort replacements by length to avoid partial replacement issues (longest first)
        sorted_reps = sorted(replacements.items(), key=lambda x: len(x[0]), reverse=True)

        for k, v in sorted_reps:
            if k not in bold_keys:
                text = text.replace(k, str(v))

        if any(k in text for k in bold_keys):
            p.clear()

            def rebuild_p(current_text, current_p):
                # Find the first bold key that appears in current_text
                first_key = None
                first_idx = len(current_text)
                for k in bold_keys:
                    idx = current_text.find(k)
                    if idx != -1 and idx < first_idx:
                        first_key = k
                        first_idx = idx

                if first_key:
                    parts = current_text.split(first_key, 1)
                    if parts[0]:
                        current_p.add_run(parts[0])
                    run = current_p.add_run(str(replacements[first_key]))
                    run.bold = True
                    rebuild_p(parts[1], current_p)
                else:
                    if current_text:
                        current_p.add_run(current_text)

            rebuild_p(text, p)
        else:
            p.text = text

advanced_replace([p], replacements, bold_keys)
print(p.text)
for run in p.runs:
    print(f"'{run.text}' bold: {run.bold}")
