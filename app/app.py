import os
import subprocess
from datetime import datetime
from flask import Flask, request, send_file, render_template
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

app = Flask(__name__)

TEMPLATES = {
    "102_options_shares": {
        "en": "Holding Approval - 102 trustee - options & shares.docx",
        "he": "אישור החזקה- 102 נאמן - מניות ואופציות.docx"
    },
    "102_shares": {
        "en": "Holding Approval - 102 trustee -Shares.docx",
        "he": "אישור החזקה- 102 נאמן - מניות.docx"
    },
    "103K_shares": {
        "en": "Holding Approval - 103K - Shares.docx",
        "he": "אישור החזקה - 103כ - מניות.docx"
    },
    "104B_shares": {
        "en": "Holding Approval - 104B - Shares.docx",
        "he": "אישור החזקה - 104ב - מניות.docx"
    },
    "104H_shares": {
        "en": "Holding Approval - 104H - Shares.docx",
        "he": "אישור החזקה - 104ח - מניות_1.docx"
    },
    "15C_shares": {
        "en": "Holding Approval - 15C - Shares.docx",
        "he": "אישור החזקה - 15ג - מניות.docx"
    },
    "3i_shares": {
        "en": "Holding Approval - 3i - Shares.docx",
        "he": "Holding Approval - 3i - Shares.docx"
    },
    "shareholder": {
        "en": "Holding Approval - Shareholder.docx",
        "he": "אישור החזקה - בעל מניות- מניות.docx"
    }
}

HEBREW_MONTHS = {
    1: 'בינואר', 2: 'בפברואר', 3: 'במרץ', 4: 'באפריל', 5: 'במאי', 6: 'ביוני',
    7: 'ביולי', 8: 'באוגוסט', 9: 'בספטמבר', 10: 'באוקטובר', 11: 'בנובמבר', 12: 'בדצמבר'
}

def format_today_date(date_str):
    if not date_str: return ""
    try:
        d = datetime.strptime(date_str, '%Y-%m-%d')
        return d.strftime('%d/%m/%y')
    except: return date_str

def format_holding_date_hebrew(date_str):
    if not date_str: return ""
    try:
        d = datetime.strptime(date_str, '%Y-%m-%d')
        return f"ה-{d.day} {HEBREW_MONTHS[d.month]} {d.year}"
    except: return date_str

@app.route("/")
def index():
    return render_template("index.html")

def get_full_text(paragraph):
    return "".join(run.text for run in paragraph.runs)

def advanced_replace(paragraphs, replacements, bold_keys):
    # Sort replacements by length descending so longer phrases match first
    sorted_reps = sorted(replacements.items(), key=lambda x: len(x[0]), reverse=True)

    for p in paragraphs:
        text = get_full_text(p)

        needs_replace = False
        for key in replacements.keys():
            if key in text:
                needs_replace = True
                break

        if needs_replace:
            # Save paragraph formatting
            alignment = p.alignment
            style = p.style

            # First pass: replace non-bold text
            for k, v in sorted_reps:
                if k not in bold_keys:
                    text = text.replace(k, str(v))

            # Second pass: check if there are any bold keys left in the text
            if any(k in text for k in bold_keys):
                # Retrieve the font size and name of the first run if possible
                first_run_font_size = None
                first_run_font_name = None
                if p.runs:
                    first_run_font_size = p.runs[0].font.size
                    first_run_font_name = p.runs[0].font.name

                p.clear()

                def rebuild_p(current_text, current_p):
                    first_key = None
                    first_idx = len(current_text)
                    for k in bold_keys:
                        idx = current_text.find(k)
                        if idx != -1 and idx < first_idx:
                            first_key = k
                            first_idx = idx

                    if first_key:
                        parts = current_text.split(first_key, 1)
                        if parts[0]:
                            run_normal = current_p.add_run(parts[0])
                            run_normal.font.size = first_run_font_size
                            run_normal.font.name = first_run_font_name

                        run_bold = current_p.add_run(str(replacements[first_key]))
                        run_bold.bold = True
                        run_bold.font.size = first_run_font_size
                        run_bold.font.name = first_run_font_name

                        rebuild_p(parts[1], current_p)
                    else:
                        if current_text:
                            run_normal = current_p.add_run(current_text)
                            run_normal.font.size = first_run_font_size
                            run_normal.font.name = first_run_font_name

                rebuild_p(text, p)
            else:
                # If no bold keys were actually found, just restore text normally
                # But keep it inside the run structure to not lose formatting
                if p.runs:
                    p.runs[0].text = text
                    for run in p.runs[1:]:
                        run.text = ""
                else:
                    p.add_run(text)

            if alignment is not None:
                p.alignment = alignment

def insert_signature(doc, signature_path):
    if not os.path.exists(signature_path):
        return

    last_company_mention = -1
    for i, p in enumerate(doc.paragraphs):
        text = get_full_text(p).strip()
        # Look for the last line that matches the sign off
        if text == "אלטשר נאמנויות בע\"מ" or text == "Altshare Trusts Ltd.":
            last_company_mention = i

    if last_company_mention != -1:
        # Insert a paragraph before the sign off for the image
        p = doc.paragraphs[last_company_mention]
        new_p = p.insert_paragraph_before("")
        new_p.alignment = p.alignment if p.alignment is not None else WD_PARAGRAPH_ALIGNMENT.RIGHT # default to right for hebrew usually, or match p
        r = new_p.add_run()
        r.add_picture(signature_path, width=Inches(1.5))

@app.route("/generate", methods=["POST"])
def generate():
    template_type = request.form.get("template_type")
    language = request.form.get("language", "he")

    company_name = request.form.get("company_name", "")
    shareholder_name = request.form.get("shareholder_name", "")
    id_number = request.form.get("id_number", "")
    today_date_raw = request.form.get("today_date", "")
    holding_date_raw = request.form.get("holding_date", "")

    number_of_shares = request.form.get("number_of_shares", "")
    total_options = request.form.get("total_options", "")
    vested_options = request.form.get("vested_options", "")
    unvested_options = request.form.get("unvested_options", "")

    today_date_formatted = format_today_date(today_date_raw)

    bold_keys = []

    if language == "en":
        replacements = {
            "[Name of Company]": company_name,
            "[company’s name]": company_name,
            "[Name of Shareholder]": shareholder_name,
            "[ID number]": id_number,
            "[DD MM, YYYY]": today_date_formatted,
            "[New Version]": "",
            "[Ordinary]": "Ordinary",
            "[Sir/Madam]": "Sir/Madam",
            "August 21, 2026": today_date_formatted,
            "July 29, 2026": today_date_formatted,
            "as of the date of this letter": f"as of {holding_date_raw}" if holding_date_raw else "as of the date of this letter",

            # English Bold Keys replacements
            "[Number of Shares] Company shares": f"{number_of_shares} Company shares",
            "[Total Number of options]": f"{total_options} options",
            "[Vested options] options are vested": f"{vested_options} options are vested",
            "[Non Vested options] options have not yet reached their vesting date": f"{unvested_options} options have not yet reached their vesting date",
        }

        # English Fallbacks
        replacements["[Number of Shares]"] = number_of_shares
        replacements["[Vested options]"] = vested_options
        replacements["[Non Vested options]"] = unvested_options

        bold_keys = [
            "[Number of Shares] Company shares",
            "[Total Number of options]",
            "[Vested options] options are vested",
            "[Non Vested options] options have not yet reached their vesting date"
        ]
    else:
        holding_date_heb = format_holding_date_hebrew(holding_date_raw)

        replacements = {
            "[שם החברה]": company_name,
            "[שם הניצע\\ת]": shareholder_name,
            "[שם הניצע\\]": shareholder_name,
            "[שם בעל המניות]": shareholder_name,
            "[תעודת זהות \\ ח\"פ]": id_number,
            "[מס' זהות/ דרכון]": id_number,
            "[מס' ת.ז.]": id_number,
            "[DD בMM, YYYY]": today_date_formatted,
            "[נא לעדכן בהתאם לאדון/גברת]": "",
            "[נוסח חדש]": "",
            "21 August 2026": today_date_formatted,
            "29 יולי 2026": today_date_formatted,
            "‏29 יולי 2026": today_date_formatted,
            "‏‏29 יולי 2026": today_date_formatted,
            "יצחק וייס": shareholder_name,
            "014875538": id_number,
            "Kanabo Group PLC": company_name,
            "נכון למועד מכתב זה": f"נכון למועד {holding_date_heb}" if holding_date_raw else "נכון למועד מכתב זה",

            # Bold Keys Replacements
            # Note: We include variations with multiple spaces that appear in the raw text
            "[מספר המניות] מניות [רגילות]": f"{number_of_shares} מניות רגילות",
            "[מספר המניות] מניות רגילות": f"{number_of_shares} מניות רגילות",
            "[כמות מניות] מניות רגילות": f"{number_of_shares} מניות רגילות",

            "[סך הכול אופציות]  מתוכן": f"{total_options} אופציות מתוכן",
            "[סך הכול אופציות] מתוכן": f"{total_options} אופציות מתוכן",

            "[אופציות מובשלות]  אופציות בשלות": f"{vested_options} אופציות בשלות",
            "[אופציות מובשלות] אופציות בשלות": f"{vested_options} אופציות בשלות",

            "[אופציות לא מובשלות]  אופציות שטרם הגיע מועד הבשלתן": f"{unvested_options} אופציות שטרם הגיע מועד הבשלתן",
            "[אופציות לא מובשלות] אופציות שטרם הגיע מועד הבשלתן": f"{unvested_options} אופציות שטרם הגיע מועד הבשלתן",
        }

        # Fallbacks
        replacements["[מספר המניות]"] = number_of_shares
        replacements["[כמות מניות]"] = number_of_shares
        replacements["[סך הכול אופציות]"] = total_options
        replacements["[אופציות מובשלות]"] = vested_options
        replacements["[אופציות לא מובשלות]"] = unvested_options
        replacements["[רגילות]"] = "רגילות"

        bold_keys = [
            "[מספר המניות] מניות [רגילות]",
            "[מספר המניות] מניות רגילות",
            "[כמות מניות] מניות רגילות",

            "[סך הכול אופציות]  מתוכן",
            "[סך הכול אופציות] מתוכן",

            "[אופציות מובשלות] אופציות בשלות",
            "[אופציות מובשלות]  אופציות בשלות",

            "[אופציות לא מובשלות] אופציות שטרם הגיע מועד הבשלתן",
            "[אופציות לא מובשלות]  אופציות שטרם הגיע מועד הבשלתן"
        ]

    template_file = TEMPLATES.get(template_type, {}).get(language)
    if not template_file:
        return "Template not found", 404

    base_dir = os.path.dirname(os.path.abspath(__file__))
    template_path = os.path.join(base_dir, "docx_templates", template_file)

    if not os.path.exists(template_path):
        return f"File {template_path} not found", 404

    doc = docx.Document(template_path)

    advanced_replace(doc.paragraphs, replacements, bold_keys)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                advanced_replace(cell.paragraphs, replacements, bold_keys)

    signature_path = os.path.join(base_dir, "static", "signature.png")
    insert_signature(doc, signature_path)

    output_docx = f"/tmp/{template_type}_{language}.docx"
    doc.save(output_docx)

    subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", "/tmp", output_docx], check=True)

    output_pdf = f"/tmp/{template_type}_{language}.pdf"

    return send_file(output_pdf, as_attachment=True)

if __name__ == "__main__":
    app.run(debug=True, host='0.0.0.0', port=5000)
